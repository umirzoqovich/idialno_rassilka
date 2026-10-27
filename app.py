import os
import sys

# Ensure UTF-8 output on Windows consoles
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import json
import asyncio
import random
import re
import datetime
import threading
import webbrowser
from pathlib import Path

from aiohttp import web, WSMsgType
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from telethon import TelegramClient, errors
from telethon.tl.functions.contacts import ImportContactsRequest, DeleteContactsRequest
from telethon.tl.types import InputPhoneContact

import database

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
UPLOADS_DIR = BASE_DIR / "uploads"
EXPORTS_DIR = BASE_DIR / "exports"
SESSION_DIR = BASE_DIR / "session"
DATA_DIR = BASE_DIR / "data"

for d in [STATIC_DIR, UPLOADS_DIR, EXPORTS_DIR, SESSION_DIR, DATA_DIR]:
    d.mkdir(parents=True, exist_ok=True)

# Standard official Telegram API credentials (fallback for seamless 1-click connect)
DEFAULT_API_ID = "2040"
DEFAULT_API_HASH = "b18441a1ff607e10a989891a5462e627"

# ----------------- WebSocket Hub (Per-User) -----------------
user_ws_connections: dict[int, set[web.WebSocketResponse]] = {}

async def broadcast_user_ws(user_id: int, event_type: str, data: dict):
    sockets = user_ws_connections.get(user_id, set()).copy()
    if not sockets:
        return
    payload = json.dumps({"type": event_type, "data": data}, ensure_ascii=False)
    dead = set()
    for ws in sockets:
        try:
            if not ws.closed:
                await ws.send_str(payload)
            else:
                dead.add(ws)
        except Exception:
            dead.add(ws)
    if dead and user_id in user_ws_connections:
        user_ws_connections[user_id].difference_update(dead)

# ----------------- Phone Number Cleaner -----------------
def clean_phone_number(val):
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        try:
            val = str(int(val))
        except (ValueError, OverflowError):
            val = str(val)
    val = str(val).strip()
    digits = re.sub(r"[^\d+]", "", val)
    if not digits:
        return ""
    # Standardize Uzbekistan or international formats
    raw_digits = digits.replace("+", "")
    if len(raw_digits) == 9 and raw_digits.startswith(("9", "8", "7", "3", "5")):
        return f"+998{raw_digits}"
    elif len(raw_digits) == 12 and raw_digits.startswith("998"):
        return f"+{raw_digits}"
    elif not digits.startswith("+"):
        return f"+{digits}"
    return digits

# ----------------- Excel Processing -----------------
def parse_excel_file(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {"error": "Excel fayl bo'sh!"}

    # Find headers (row 0 or first non-empty row)
    header_idx = 0
    headers = []
    for idx, r in enumerate(rows):
        if any(cell is not None and str(cell).strip() != "" for cell in r):
            headers = [str(c).strip() if c is not None else f"Ustun {i+1}" for i, c in enumerate(r)]
            header_idx = idx
            break

    if not headers:
        return {"error": "Excel ustunlari topilmadi!"}

    data_rows = rows[header_idx + 1:]
    
    # Auto-detect phone and name column
    phone_col_idx = None
    name_col_idx = None

    phone_keywords = ["telefon", "phone", "tel", "nomer", "raqam", "mobile", "contact", "aloqa"]
    name_keywords = ["ism", "name", "fio", "mijoz", "nomi", "first_name", "user", "клиент", "имя"]

    for i, h in enumerate(headers):
        h_lower = h.lower()
        if phone_col_idx is None and any(k in h_lower for k in phone_keywords):
            phone_col_idx = i
        if name_col_idx is None and any(k in h_lower for k in name_keywords):
            name_col_idx = i

    # Inspect sample data if not detected by header name
    if phone_col_idx is None:
        for i in range(len(headers)):
            sample_matches = 0
            for r in data_rows[:15]:
                if i < len(r) and r[i]:
                    cleaned = clean_phone_number(r[i])
                    if len(cleaned.replace("+", "")) >= 9:
                        sample_matches += 1
            if sample_matches >= 2:
                phone_col_idx = i
                break

    if phone_col_idx is None:
        phone_col_idx = 0  # default to first column

    contacts = []
    preview_rows = []

    for r_idx, r in enumerate(data_rows):
        if not r or all(c is None or str(c).strip() == "" for c in r):
            continue
        
        raw_phone = r[phone_col_idx] if phone_col_idx < len(r) else ""
        cleaned_phone = clean_phone_number(raw_phone)
        if not cleaned_phone:
            continue

        raw_name = ""
        if name_col_idx is not None and name_col_idx < len(r):
            raw_name = str(r[name_col_idx]).strip() if r[name_col_idx] is not None else ""

        # Map all other columns for possible variables
        row_dict = {}
        for c_idx, val in enumerate(r):
            col_name = headers[c_idx] if c_idx < len(headers) else f"col_{c_idx+1}"
            row_dict[col_name] = str(val).strip() if val is not None else ""

        contact_entry = {
            "id": len(contacts) + 1,
            "raw_phone": str(raw_phone),
            "phone": cleaned_phone,
            "name": raw_name,
            "extra": row_dict
        }
        contacts.append(contact_entry)
        if len(preview_rows) < 10:
            preview_rows.append({
                "id": contact_entry["id"],
                "phone": contact_entry["phone"],
                "name": contact_entry["name"] or "—",
                "raw": [str(c) if c is not None else "" for c in r[:len(headers)]]
            })

    return {
        "headers": headers,
        "phone_col_idx": phone_col_idx,
        "name_col_idx": name_col_idx,
        "total_contacts": len(contacts),
        "preview": preview_rows,
        "contacts": contacts
    }

# ----------------- Per-User Telegram Manager -----------------
class UserTelegramManager:
    def __init__(self, user_id: int):
        self.user_id = user_id
        self.client = None
        self.api_id = None
        self.api_hash = None
        self.phone = None
        self.phone_code_hash = None
        self.session_path = str(SESSION_DIR / f"user_{user_id}")
        self.is_connected = False
        self.is_authorized = False
        self.user_info = None

    async def init_client(self, api_id, api_hash):
        if self.client:
            try:
                await self.client.disconnect()
            except Exception:
                pass
            self.client = None

        self.api_id = int(api_id)
        self.api_hash = str(api_hash).strip()
        self.client = TelegramClient(self.session_path, self.api_id, self.api_hash)
        await self.client.connect()
        self.is_connected = True
        self.is_authorized = await self.client.is_user_authorized()
        if self.is_authorized:
            me = await self.client.get_me()
            self.user_info = {
                "id": me.id,
                "first_name": me.first_name or "",
                "last_name": me.last_name or "",
                "username": me.username or "",
                "phone": me.phone or ""
            }
        else:
            self.user_info = None
        return self.is_authorized

    async def check_auth(self):
        if self.client and self.client.is_connected():
            self.is_authorized = await self.client.is_user_authorized()
            if self.is_authorized:
                me = await self.client.get_me()
                self.user_info = {
                    "id": me.id,
                    "first_name": me.first_name or "",
                    "last_name": me.last_name or "",
                    "username": me.username or "",
                    "phone": me.phone or ""
                }
            return self.is_authorized
        return False

    async def request_code(self, phone):
        if not self.client or not self.client.is_connected():
            raise Exception("Telegram mijoz ulangan emas!")
        self.phone = phone.strip()
        res = await self.client.send_code_request(self.phone)
        self.phone_code_hash = res.phone_code_hash
        return {"phone_code_hash": self.phone_code_hash, "phone": self.phone}

    async def sign_in_with_code(self, code):
        if not self.client:
            raise Exception("Telegram mijoz ulanmagan!")
        try:
            await self.client.sign_in(phone=self.phone, code=code.strip(), phone_code_hash=self.phone_code_hash)
            self.is_authorized = True
            me = await self.client.get_me()
            self.user_info = {
                "id": me.id,
                "first_name": me.first_name or "",
                "last_name": me.last_name or "",
                "username": me.username or "",
                "phone": me.phone or ""
            }
            return {"status": "success", "user": self.user_info}
        except errors.SessionPasswordNeededError:
            return {"status": "password_needed"}

    async def sign_in_with_password(self, password):
        if not self.client:
            raise Exception("Telegram mijoz ulanmagan!")
        await self.client.sign_in(password=password)
        self.is_authorized = True
        me = await self.client.get_me()
        self.user_info = {
            "id": me.id,
            "first_name": me.first_name or "",
            "last_name": me.last_name or "",
            "username": me.username or "",
            "phone": me.phone or ""
        }
        return {"status": "success", "user": self.user_info}

    async def logout(self):
        if self.client:
            try:
                await self.client.log_out()
            except Exception:
                pass
            try:
                await self.client.disconnect()
            except Exception:
                pass
            self.client = None
        self.is_connected = False
        self.is_authorized = False
        self.user_info = None
        # remove user session files
        for f in SESSION_DIR.glob(f"user_{self.user_id}.*"):
            try:
                f.unlink()
            except Exception:
                pass
        return True

telegram_managers: dict[int, UserTelegramManager] = {}

def get_user_telegram_mgr(user_id: int) -> UserTelegramManager:
    if user_id not in telegram_managers:
        telegram_managers[user_id] = UserTelegramManager(user_id)
    return telegram_managers[user_id]

# ----------------- Per-User Campaign Manager -----------------
class UserCampaignManager:
    def __init__(self, user_id: int):
        self.user_id = user_id
        self.status = "idle"  # idle, running, paused, stopped, completed
        self.task = None
        self.pause_event = asyncio.Event()
        self.pause_event.set()
        
        self.contacts = []
        self.message_template = ""
        self.media_path = None
        self.min_delay = 15
        self.max_delay = 30
        self.auto_clean_contacts = True
        self.max_limit = 0
        
        self.stats = {
            "total": 0,
            "processed": 0,
            "sent": 0,
            "not_found": 0,
            "failed": 0,
            "current_status_text": "Kutilmoqda",
            "start_time": None,
            "end_time": None
        }
        self.logs = []
        self.results = []
        self.last_export_file = None
        self.last_export_filename = None

    def reset(self):
        self.status = "idle"
        self.task = None
        self.pause_event.set()
        self.contacts = []
        self.message_template = ""
        self.media_path = None
        self.stats = {
            "total": 0,
            "processed": 0,
            "sent": 0,
            "not_found": 0,
            "failed": 0,
            "current_status_text": "Kutilmoqda",
            "start_time": None,
            "end_time": None
        }
        self.logs = []
        self.results = []

    def add_log(self, phone, name, status, message):
        log_entry = {
            "id": len(self.logs) + 1,
            "time": datetime.datetime.now().strftime("%H:%M:%S"),
            "phone": phone,
            "name": name or "—",
            "status": status,  # sent, not_found, failed, info, warning
            "message": message
        }
        self.logs.append(log_entry)
        # Asynchronously send log to user's WebSocket
        asyncio.create_task(broadcast_user_ws(self.user_id, "new_log", {"log": log_entry, "stats": self.stats}))

    async def run(self):
        tg_mgr = get_user_telegram_mgr(self.user_id)
        if not tg_mgr.client or not tg_mgr.is_authorized:
            self.status = "stopped"
            self.add_log("Tizim", "Tizim", "failed", "Telegram akkaunt ulanmagan!")
            await broadcast_user_ws(self.user_id, "status_change", {"status": self.status, "stats": self.stats})
            return

        self.status = "running"
        self.stats["start_time"] = datetime.datetime.now().isoformat()
        
        to_process = self.contacts
        if self.max_limit > 0 and self.max_limit < len(to_process):
            to_process = to_process[:self.max_limit]

        self.stats["total"] = len(to_process)
        self.stats["processed"] = 0
        self.stats["sent"] = 0
        self.stats["not_found"] = 0
        self.stats["failed"] = 0
        self.stats["current_status_text"] = "Rassilka boshlandi..."

        await broadcast_user_ws(self.user_id, "status_change", {"status": "running", "stats": self.stats})
        self.add_log("Tizim", "Tizim", "info", f"Rassilka boshlandi. Jami: {len(to_process)} ta mijoz.")

        for idx, contact in enumerate(to_process):
            if self.status == "stopped":
                break

            await self.pause_event.wait()
            if self.status == "stopped":
                break

            phone = contact["phone"]
            name = contact.get("name", "")
            extra = contact.get("extra", {})

            # Prepare personalized message
            msg = self.message_template
            msg = msg.replace("{ism}", name or "").replace("{name}", name or "")
            msg = msg.replace("{telefon}", phone).replace("{phone}", phone)
            for k, v in extra.items():
                msg = msg.replace(f"{{{k}}}", str(v))

            client = tg_mgr.client
            if not client or not tg_mgr.is_authorized:
                self.add_log("Tizim", "Tizim", "failed", "Telegram akkaunt uzildi!")
                self.status = "stopped"
                break

            result_entry = {
                "phone": phone,
                "name": name,
                "status": "failed",
                "detail": "",
                "time": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }

            try:
                first_n = name if name else "Mijoz"
                input_contact = InputPhoneContact(client_id=random.randint(1000, 99999), phone=phone, first_name=first_n, last_name="")
                import_res = await client(ImportContactsRequest([input_contact]))

                if not import_res.users:
                    self.stats["not_found"] += 1
                    result_entry["status"] = "not_found"
                    result_entry["detail"] = "Telegramda ro'yxatdan o'tmagan"
                    self.add_log(phone, name, "not_found", "Telegram topilmadi")
                else:
                    target_user = import_res.users[0]
                    if self.media_path and os.path.exists(self.media_path):
                        await client.send_file(target_user, self.media_path, caption=msg)
                    else:
                        await client.send_message(target_user, msg)

                    self.stats["sent"] += 1
                    result_entry["status"] = "sent"
                    result_entry["detail"] = "Xabar muvaffaqiyatli yuborildi"
                    self.add_log(phone, name, "sent", "Xabar yuborildi")

                    if self.auto_clean_contacts:
                        try:
                            await client(DeleteContactsRequest(id=[target_user.id]))
                        except Exception:
                            pass

            except errors.FloodWaitError as e:
                wait_sec = e.seconds
                self.add_log("Telegram", "Limit", "warning", f"Telegram FloodWait cheklovi! {wait_sec} soniya kutilmoqda...")
                self.stats["current_status_text"] = f"Telegram cheklovi: {wait_sec}s kutilmoqda"
                await broadcast_user_ws(self.user_id, "progress", {"stats": self.stats})
                await asyncio.sleep(wait_sec + 2)
                try:
                    input_contact = InputPhoneContact(client_id=random.randint(1000, 99999), phone=phone, first_name=name or "Mijoz", last_name="")
                    import_res = await client(ImportContactsRequest([input_contact]))
                    if import_res.users:
                        target_user = import_res.users[0]
                        if self.media_path and os.path.exists(self.media_path):
                            await client.send_file(target_user, self.media_path, caption=msg)
                        else:
                            await client.send_message(target_user, msg)
                        self.stats["sent"] += 1
                        result_entry["status"] = "sent"
                        result_entry["detail"] = "Qayta urinishda yuborildi"
                        self.add_log(phone, name, "sent", "Qayta urinishda yuborildi")
                    else:
                        self.stats["not_found"] += 1
                        result_entry["status"] = "not_found"
                        result_entry["detail"] = "Telegram topilmadi"
                        self.add_log(phone, name, "not_found", "Telegram topilmadi")
                except Exception as retry_err:
                    self.stats["failed"] += 1
                    result_entry["status"] = "failed"
                    result_entry["detail"] = str(retry_err)
                    self.add_log(phone, name, "failed", f"Xato: {retry_err}")

            except errors.PeerFloodError:
                self.stats["failed"] += 1
                result_entry["status"] = "failed"
                result_entry["detail"] = "Telegram SpamBlock (PeerFloodError). Xabar yuborish to'xtatildi."
                self.add_log("Telegram", "SpamBlock", "failed", "PeerFloodError: Telegram akkaunt cheklandi! Rassilka to'xtatilmoqda.")
                self.results.append(result_entry)
                self.status = "stopped"
                break

            except errors.UserPrivacyRestrictedError:
                self.stats["failed"] += 1
                result_entry["status"] = "failed"
                result_entry["detail"] = "Foydalanuvchi maxfiylik sozlamalari (xabar qabul qilmaydi)"
                self.add_log(phone, name, "failed", "Maxfiylik cheklovi")

            except Exception as err:
                self.stats["failed"] += 1
                result_entry["status"] = "failed"
                result_entry["detail"] = str(err)
                self.add_log(phone, name, "failed", f"Xatolik: {str(err)[:50]}")

            self.results.append(result_entry)
            self.stats["processed"] = idx + 1
            await broadcast_user_ws(self.user_id, "progress", {"stats": self.stats})

            # Delay before next message if not last item
            if idx < len(to_process) - 1 and self.status == "running":
                sleep_sec = random.randint(min(self.min_delay, self.max_delay), max(self.min_delay, self.max_delay))
                self.stats["current_status_text"] = f"Keyingi xabargacha: {sleep_sec} soniya kutilmoqda..."
                await broadcast_user_ws(self.user_id, "progress", {"stats": self.stats})
                
                for _ in range(sleep_sec):
                    if self.status == "stopped" or not self.pause_event.is_set():
                        break
                    await asyncio.sleep(1)

        if self.status != "stopped":
            self.status = "completed"
            self.stats["current_status_text"] = "Rassilka muvaffaqiyatli yakunlandi!"
            self.add_log("Tizim", "Tizim", "info", f"Rassilka yakunlandi! Jami yuborildi: {self.stats['sent']} ta.")
        
        self.stats["end_time"] = datetime.datetime.now().isoformat()
        self.stats["status"] = self.status

        # Auto-export results and save to history
        try:
            exp_path, exp_name = export_results_to_excel(self.results, self.user_id)
            self.last_export_file = str(exp_path)
            self.last_export_filename = exp_name
            await database.save_campaign_history(
                user_id=self.user_id,
                stats=self.stats,
                message_preview=self.message_template,
                report_filename=exp_name,
                report_path=str(exp_path),
                results=self.results
            )
        except Exception as e:
            print(f"User {self.user_id} campaign history save error: {e}")

        await broadcast_user_ws(self.user_id, "status_change", {"status": self.status, "stats": self.stats})

campaign_managers: dict[int, UserCampaignManager] = {}

def get_user_campaign_mgr(user_id: int) -> UserCampaignManager:
    if user_id not in campaign_managers:
        campaign_managers[user_id] = UserCampaignManager(user_id)
    return campaign_managers[user_id]

# ----------------- Excel Results Exporter -----------------
def export_results_to_excel(results, user_id: int):
    user_export_dir = EXPORTS_DIR / str(user_id)
    user_export_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rassilka Natijalari"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = ["№", "Telefon Raqam", "Ism", "Holati", "Tafsilot", "Vaqt"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    status_colors = {
        "sent": ("D1FAE5", "065F46", "Yuborildi"),
        "not_found": ("FEF3C7", "92400E", "Telegram topilmadi"),
        "failed": ("FEE2E2", "991B1B", "Xatolik")
    }

    for idx, r in enumerate(results, start=1):
        status_key = r.get("status", "failed")
        bg_col, txt_col, label = status_colors.get(status_key, ("F1F5F9", "334155", status_key))
        
        row_vals = [
            idx,
            r.get("phone", ""),
            r.get("name", "") or "—",
            label,
            r.get("detail", ""),
            r.get("time", "")
        ]
        ws.append(row_vals)
        curr_row = idx + 1
        
        for c in range(1, len(row_vals) + 1):
            cell = ws.cell(row=curr_row, column=c)
            cell.border = thin_border
            cell.alignment = align_center if c in (1, 4, 6) else align_left
            if c == 4:
                cell.fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
                cell.font = Font(name="Segoe UI", bold=True, color=txt_col)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    file_name = f"rassilka_natijalari_{timestamp}.xlsx"
    export_path = user_export_dir / file_name
    wb.save(export_path)
    return export_path, file_name

# ----------------- Sample Template Generator -----------------
def generate_sample_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Namunaviy Kontaktlar"

    headers = ["Telefon raqam", "Ism", "Kompaniya", "Qo'shimcha"]
    sample_data = [
        ["+998901234567", "Azizbek", "TechCorp", "Mijoz"],
        ["+998939876543", "Madina", "GrandSoft", "Hamkor"],
        ["+998971112233", "Jasur", "EcoStroy", "Mijoz"],
        ["+998994445566", "Dilshod", "AutoTrade", "VIP"],
        ["+998947778899", "Shahnoza", "BeautyLab", "Yangi"]
    ]

    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for r in sample_data:
        ws.append(r)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

    sample_path = EXPORTS_DIR / "namuna_kontaktlar.xlsx"
    wb.save(sample_path)
    return sample_path

# ----------------- HTTP Middlewares & Handlers -----------------
@web.middleware
async def auth_middleware(request, handler):
    token = request.cookies.get("auth_token")
    if not token:
        auth_header = request.headers.get("Authorization")
        if auth_header and auth_header.startswith("Bearer "):
            token = auth_header[7:].strip()
    if not token and request.query.get("token"):
        token = request.query.get("token")

    user = None
    if token:
        user = await database.get_user_by_session(token)

    request["user"] = user
    request["user_id"] = user["id"] if user else None

    # Public API paths
    public_paths = [
        "/api/auth/login",
        "/api/auth/register",
        "/api/auth/google",
        "/api/auth/firebase-config",
        "/api/excel/template",
    ]

    if request.path.startswith("/api/") and request.path not in public_paths:
        if not user:
            return web.json_response({
                "ok": False,
                "error": "Iltimos, avval tizimga kiring!",
                "unauthorized": True
            }, status=401)

    return await handler(request)

async def handle_index(request):
    return web.FileResponse(STATIC_DIR / "index.html")

# ----------------- User Auth Handlers -----------------
async def handle_user_register(request):
    data = await request.json()
    email = data.get("email", "").strip()
    password = data.get("password", "").strip()
    full_name = data.get("full_name", "").strip()

    if not email:
        return web.json_response({"ok": False, "error": "Gmail / Pochta manzilingizni kiriting!"}, status=400)
    if "@" not in email or "." not in email:
        return web.json_response({"ok": False, "error": "Yaroqli Gmail yoki elektron pochta kiriting!"}, status=400)
    if len(password) < 6:
        return web.json_response({"ok": False, "error": "Parol kamida 6 ta belgidan iborat bo'lishi kerak!"}, status=400)

    try:
        user = await database.create_user(email, password, full_name)
        token = await database.create_session(user["id"])
        
        resp = web.json_response({
            "ok": True,
            "user": user,
            "token": token,
            "message": "Muvaffaqiyatli ro'yxatdan o'tdingiz!"
        })
        resp.set_cookie("auth_token", token, max_age=30*86400, httponly=True, samesite="Lax", path="/")
        return resp
    except ValueError as ve:
        return web.json_response({"ok": False, "error": str(ve)}, status=400)
    except Exception as e:
        return web.json_response({"ok": False, "error": f"Tizim xatoligi: {str(e)}"}, status=500)

async def handle_user_login(request):
    data = await request.json()
    email = data.get("email", "").strip()
    password = data.get("password", "").strip()

    if not email or not password:
        return web.json_response({"ok": False, "error": "Gmail va parolni to'liq kiriting!"}, status=400)

    user = await database.authenticate_user(email, password)
    if not user:
        return web.json_response({"ok": False, "error": "Pochta manzili yoki parol noto'g'ri!"}, status=401)

    token = await database.create_session(user["id"])
    
    # Auto-connect Telegram if user already saved credentials
    user_settings = await database.get_user_settings(user["id"])
    if user_settings.get("api_id") and user_settings.get("api_hash"):
        tg_mgr = get_user_telegram_mgr(user["id"])
        try:
            await tg_mgr.init_client(user_settings["api_id"], user_settings["api_hash"])
        except Exception:
            pass

    resp = web.json_response({
        "ok": True,
        "user": {
            "id": user["id"],
            "email": user["email"],
            "full_name": user["full_name"],
            "avatar": user["avatar"]
        },
        "token": token
    })
    resp.set_cookie("auth_token", token, max_age=30*86400, httponly=True, samesite="Lax", path="/")
    return resp

async def handle_google_auth(request):
    data = await request.json()
    email = data.get("email", "").strip()
    full_name = data.get("full_name", "").strip()
    avatar = data.get("avatar", "").strip()
    credential = data.get("credential", "").strip()

    if credential and not email:
        try:
            import aiohttp
            async with aiohttp.ClientSession() as http_sess:
                async with http_sess.get(f"https://oauth2.googleapis.com/tokeninfo?id_token={credential}") as g_resp:
                    if g_resp.status == 200:
                        g_info = await g_resp.json()
                        email = g_info.get("email", "").strip()
                        full_name = g_info.get("name", "").strip()
                        avatar = g_info.get("picture", "").strip()
        except Exception as e:
            print(f"Google token verification error: {e}")

    if not email:
        return web.json_response({"ok": False, "error": "Gmail manzili topilmadi!"}, status=400)

    user = await database.get_or_create_google_user(email, full_name, avatar)
    token = await database.create_session(user["id"])

    # Auto-connect Telegram if user settings exist
    user_settings = await database.get_user_settings(user["id"])
    if user_settings.get("api_id") and user_settings.get("api_hash"):
        tg_mgr = get_user_telegram_mgr(user["id"])
        try:
            await tg_mgr.init_client(user_settings["api_id"], user_settings["api_hash"])
        except Exception:
            pass

    resp = web.json_response({
        "ok": True,
        "user": {
            "id": user["id"],
            "email": user["email"],
            "full_name": user["full_name"],
            "avatar": user["avatar"]
        },
        "token": token
    })
    resp.set_cookie("auth_token", token, max_age=30*86400, httponly=True, samesite="Lax", path="/")
    return resp

async def handle_user_logout(request):
    token = request.cookies.get("auth_token")
    if token:
        await database.delete_session(token)
    resp = web.json_response({"ok": True})
    resp.del_cookie("auth_token", path="/")
    return resp

async def handle_user_me(request):
    user = request.get("user")
    if not user:
        return web.json_response({"ok": False, "error": "Avtorizatsiyadan o'tilmagan"}, status=401)
    return web.json_response({
        "ok": True,
        "user": {
            "id": user["id"],
            "email": user["email"],
            "full_name": user["full_name"],
            "avatar": user["avatar"]
        }
    })

FIREBASE_CONFIG_PATH = DATA_DIR / "firebase_config.json"

async def handle_get_firebase_config(request):
    config = None
    if FIREBASE_CONFIG_PATH.exists():
        try:
            with open(FIREBASE_CONFIG_PATH, "r", encoding="utf-8") as f:
                config = json.load(f)
        except Exception:
            pass
    if not config and os.environ.get("FIREBASE_CONFIG"):
        try:
            config = json.loads(os.environ["FIREBASE_CONFIG"])
        except Exception:
            pass
    return web.json_response({"ok": True, "config": config})

async def handle_save_firebase_config(request):
    try:
        data = await request.json()
        config = data.get("config", {})
        if not isinstance(config, dict) or not config.get("apiKey"):
            return web.json_response({"ok": False, "error": "apiKey topilmadi yoki noto'g'ri Firebase config!"}, status=400)

        FIREBASE_CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
        with open(FIREBASE_CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(config, f, indent=2)

        return web.json_response({"ok": True, "message": "Firebase sozlamalari muvaffaqiyatli saqlandi!"})
    except Exception as e:
        return web.json_response({"ok": False, "error": f"Xatolik: {str(e)}"}, status=500)


# ----------------- Status & Settings (Per User) -----------------
async def handle_get_status(request):
    user_id = request["user_id"]
    user_settings = await database.get_user_settings(user_id)
    tg_mgr = get_user_telegram_mgr(user_id)
    cmp_mgr = get_user_campaign_mgr(user_id)

    # Check connection if settings exist
    if user_settings.get("api_id") and user_settings.get("api_hash") and not tg_mgr.is_connected:
        try:
            await tg_mgr.init_client(user_settings["api_id"], user_settings["api_hash"])
        except Exception:
            pass

    stats = await database.get_user_stats(user_id)
    return web.json_response({
        "user": request["user"],
        "analytics": stats,
        "telegram": {
            "is_connected": tg_mgr.is_connected,
            "is_authorized": tg_mgr.is_authorized,
            "user": tg_mgr.user_info,
            "has_credentials": bool(user_settings.get("api_id") and user_settings.get("api_hash")),
            "is_default_api": bool(not user_settings.get("api_id") or user_settings.get("api_id") == DEFAULT_API_ID)
        },
        "campaign": {
            "status": cmp_mgr.status,
            "stats": cmp_mgr.stats
        },
        "config": user_settings
    })

async def handle_save_config(request):
    user_id = request["user_id"]
    data = await request.json()
    saved = await database.save_user_settings(user_id, data)
    
    tg_mgr = get_user_telegram_mgr(user_id)
    if saved.get("api_id") and saved.get("api_hash") and not tg_mgr.is_authorized:
        try:
            await tg_mgr.init_client(saved["api_id"], saved["api_hash"])
        except Exception as e:
            return web.json_response({"ok": False, "error": str(e)})

    return web.json_response({
        "ok": True,
        "config": saved,
        "is_authorized": tg_mgr.is_authorized,
        "user": tg_mgr.user_info
    })

# ----------------- Telegram Auth (Per User) -----------------
async def handle_auth_send_code(request):
    user_id = request["user_id"]
    data = await request.json()
    api_id = str(data.get("api_id", "")).strip() or DEFAULT_API_ID
    api_hash = str(data.get("api_hash", "")).strip() or DEFAULT_API_HASH
    phone = str(data.get("phone", "")).strip()

    if not phone:
        return web.json_response({"ok": False, "error": "Telegram telefon raqamingizni kiriting!"}, status=400)

    await database.save_user_settings(user_id, {"api_id": api_id, "api_hash": api_hash, "phone": phone})
    tg_mgr = get_user_telegram_mgr(user_id)

    try:
        await tg_mgr.init_client(api_id, api_hash)
        if tg_mgr.is_authorized:
            return web.json_response({"ok": True, "status": "already_authorized", "user": tg_mgr.user_info})

        res = await tg_mgr.request_code(phone)
        return web.json_response({"ok": True, "status": "code_sent", "phone": res["phone"]})
    except Exception as e:
        return web.json_response({"ok": False, "error": f"Xatolik: {str(e)}"}, status=500)

async def handle_auth_verify_code(request):
    user_id = request["user_id"]
    data = await request.json()
    code = data.get("code")
    if not code:
        return web.json_response({"ok": False, "error": "SMS yoki Telegram kodini kiriting!"}, status=400)

    tg_mgr = get_user_telegram_mgr(user_id)
    try:
        res = await tg_mgr.sign_in_with_code(code)
        if res["status"] == "success":
            return web.json_response({"ok": True, "status": "success", "user": res["user"]})
        elif res["status"] == "password_needed":
            return web.json_response({"ok": True, "status": "password_needed"})
    except Exception as e:
        return web.json_response({"ok": False, "error": f"Tasdiqlashda xatolik: {str(e)}"}, status=400)

async def handle_auth_verify_password(request):
    user_id = request["user_id"]
    data = await request.json()
    password = data.get("password")
    if not password:
        return web.json_response({"ok": False, "error": "2FA parolini kiriting!"}, status=400)

    tg_mgr = get_user_telegram_mgr(user_id)
    try:
        res = await tg_mgr.sign_in_with_password(password)
        return web.json_response({"ok": True, "status": "success", "user": res["user"]})
    except Exception as e:
        return web.json_response({"ok": False, "error": f"Parol noto'g'ri: {str(e)}"}, status=400)

async def handle_auth_logout_telegram(request):
    user_id = request["user_id"]
    tg_mgr = get_user_telegram_mgr(user_id)
    await tg_mgr.logout()
    return web.json_response({"ok": True})

# ----------------- File Uploads (Per User) -----------------
async def handle_upload_excel(request):
    user_id = request["user_id"]
    user_upload_dir = UPLOADS_DIR / str(user_id)
    user_upload_dir.mkdir(parents=True, exist_ok=True)

    reader = await request.multipart()
    saved_path = None
    while True:
        part = await reader.next()
        if part is None:
            break
        if part.name == "excel_file":
            filename = f"contacts_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            saved_path = user_upload_dir / filename
            with open(saved_path, "wb") as f:
                while True:
                    chunk = await part.read_chunk()
                    if not chunk:
                        break
                    f.write(chunk)

    if not saved_path or not saved_path.exists():
        return web.json_response({"ok": False, "error": "Fayl yuklanmadi!"}, status=400)

    try:
        parsed = parse_excel_file(saved_path)
        if "error" in parsed:
            return web.json_response({"ok": False, "error": parsed["error"]}, status=400)
        
        parsed["file_path"] = str(saved_path)
        return web.json_response({"ok": True, "data": parsed})
    except Exception as e:
        return web.json_response({"ok": False, "error": f"Excelni o'qishda xatolik: {str(e)}"}, status=500)

async def handle_upload_media(request):
    user_id = request["user_id"]
    user_upload_dir = UPLOADS_DIR / str(user_id)
    user_upload_dir.mkdir(parents=True, exist_ok=True)

    reader = await request.multipart()
    saved_path = None
    filename = ""
    while True:
        part = await reader.next()
        if part is None:
            break
        if part.name == "media_file":
            ext = Path(part.filename).suffix if part.filename else ".jpg"
            filename = f"media_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
            saved_path = user_upload_dir / filename
            with open(saved_path, "wb") as f:
                while True:
                    chunk = await part.read_chunk()
                    if not chunk:
                        break
                    f.write(chunk)
    if saved_path and saved_path.exists():
        return web.json_response({"ok": True, "file_path": str(saved_path), "filename": filename})
    return web.json_response({"ok": False, "error": "Media yuklanmadi"}, status=400)

# ----------------- Campaign Controls (Per User) -----------------
async def handle_campaign_start(request):
    user_id = request["user_id"]
    tg_mgr = get_user_telegram_mgr(user_id)
    cmp_mgr = get_user_campaign_mgr(user_id)

    if not tg_mgr.is_authorized:
        return web.json_response({"ok": False, "error": "Avval Telegram akkauntga ulaning!"}, status=400)

    if cmp_mgr.status in ("running", "paused"):
        return web.json_response({"ok": False, "error": "Rassilka allaqachon ishlayapti!"}, status=400)

    data = await request.json()
    contacts = data.get("contacts", [])
    message = data.get("message", "").strip()
    media_path = data.get("media_path")
    min_delay = int(data.get("min_delay", 15))
    max_delay = int(data.get("max_delay", 30))
    auto_clean = bool(data.get("auto_clean_contacts", True))
    max_limit = int(data.get("max_limit", 0))

    if not contacts:
        return web.json_response({"ok": False, "error": "Hech qanday kontakt tanlanmagan!"}, status=400)
    if not message and not media_path:
        return web.json_response({"ok": False, "error": "Xabar matni yoki rasm kiritilishi shart!"}, status=400)

    cmp_mgr.reset()
    cmp_mgr.contacts = contacts
    cmp_mgr.message_template = message
    cmp_mgr.media_path = media_path
    cmp_mgr.min_delay = min_delay
    cmp_mgr.max_delay = max_delay
    cmp_mgr.auto_clean_contacts = auto_clean
    cmp_mgr.max_limit = max_limit

    # Save user preferences to DB
    await database.save_user_settings(user_id, {
        "min_delay": min_delay,
        "max_delay": max_delay,
        "auto_clean_contacts": auto_clean,
        "max_limit": max_limit
    })

    # Start user's async campaign task
    cmp_mgr.task = asyncio.create_task(cmp_mgr.run())

    return web.json_response({"ok": True, "status": "started", "total": len(contacts)})

async def handle_campaign_pause(request):
    user_id = request["user_id"]
    cmp_mgr = get_user_campaign_mgr(user_id)
    if cmp_mgr.status == "running":
        cmp_mgr.status = "paused"
        cmp_mgr.pause_event.clear()
        cmp_mgr.stats["current_status_text"] = "Vaqtincha to'xtatildi (Pauza)"
        await broadcast_user_ws(user_id, "status_change", {"status": "paused", "stats": cmp_mgr.stats})
        return web.json_response({"ok": True, "status": "paused"})
    return web.json_response({"ok": False, "error": "Rassilka faol emas"})

async def handle_campaign_resume(request):
    user_id = request["user_id"]
    cmp_mgr = get_user_campaign_mgr(user_id)
    if cmp_mgr.status == "paused":
        cmp_mgr.status = "running"
        cmp_mgr.pause_event.set()
        cmp_mgr.stats["current_status_text"] = "Yuborilmoqda..."
        await broadcast_user_ws(user_id, "status_change", {"status": "running", "stats": cmp_mgr.stats})
        return web.json_response({"ok": True, "status": "running"})
    return web.json_response({"ok": False, "error": "Rassilka pauzada emas"})

async def handle_campaign_stop(request):
    user_id = request["user_id"]
    cmp_mgr = get_user_campaign_mgr(user_id)
    if cmp_mgr.status in ("running", "paused"):
        cmp_mgr.status = "stopped"
        cmp_mgr.pause_event.set()
        cmp_mgr.stats["current_status_text"] = "To'xtatildi"
        await broadcast_user_ws(user_id, "status_change", {"status": "stopped", "stats": cmp_mgr.stats})
        return web.json_response({"ok": True, "status": "stopped"})
    return web.json_response({"ok": False, "error": "Rassilka faol emas"})

async def handle_campaign_logs(request):
    user_id = request["user_id"]
    cmp_mgr = get_user_campaign_mgr(user_id)
    return web.json_response({
        "status": cmp_mgr.status,
        "stats": cmp_mgr.stats,
        "logs": cmp_mgr.logs[-100:]
    })

async def handle_export_excel(request):
    user_id = request["user_id"]
    cmp_mgr = get_user_campaign_mgr(user_id)
    if not cmp_mgr.results:
        return web.Response(text="Hali natijalar mavjud emas!", status=404)
    file_path, file_name = export_results_to_excel(cmp_mgr.results, user_id)
    return web.FileResponse(file_path, headers={
        "Content-Disposition": f'attachment; filename="{file_name}"'
    })

async def handle_download_template(request):
    template_path = generate_sample_template()
    return web.FileResponse(template_path, headers={
        "Content-Disposition": 'attachment; filename="namuna_kontaktlar.xlsx"'
    })

# ----------------- Campaign History (Per User) -----------------
async def handle_campaign_history(request):
    user_id = request["user_id"]
    history = await database.get_user_campaign_history(user_id)
    return web.json_response({"ok": True, "history": history})

async def handle_download_history_report(request):
    user_id = request["user_id"]
    campaign_id = int(request.match_info.get("id", 0))
    campaign = await database.get_campaign_by_id(campaign_id, user_id)
    if not campaign or not campaign.get("report_path"):
        return web.Response(text="Hisobot fayli topilmadi!", status=404)
    
    file_path = Path(campaign["report_path"])
    if not file_path.exists():
        return web.Response(text="Hisobot fayli serverda topilmadi!", status=404)

    return web.FileResponse(file_path, headers={
        "Content-Disposition": f'attachment; filename="{campaign.get("report_filename", "hisobot.xlsx")}"'
    })

# ----------------- WebSocket (Per User) -----------------
async def handle_ws(request):
    user_id = request.get("user_id")
    if not user_id:
        return web.Response(status=401, text="Avtorizatsiya talab qilinadi")

    ws = web.WebSocketResponse()
    await ws.prepare(request)

    if user_id not in user_ws_connections:
        user_ws_connections[user_id] = set()
    user_ws_connections[user_id].add(ws)

    cmp_mgr = get_user_campaign_mgr(user_id)

    # Send initial state to this specific connected user
    await ws.send_str(json.dumps({
        "type": "init",
        "data": {
            "status": cmp_mgr.status,
            "stats": cmp_mgr.stats,
            "logs": cmp_mgr.logs[-50:]
        }
    }, ensure_ascii=False))

    try:
        async for msg in ws:
            if msg.type == WSMsgType.TEXT:
                pass
            elif msg.type == WSMsgType.ERROR:
                break
    finally:
        if user_id in user_ws_connections:
            user_ws_connections[user_id].discard(ws)

    return ws

# ----------------- App Lifecycle -----------------
async def on_startup(app):
    await database.init_db()
    print("[*] SQLite ma'lumotlar bazasi ishga tushirildi.")

async def on_cleanup(app):
    print("[*] Telegram mijozlari uzilmoqda...")
    for uid, mgr in list(telegram_managers.items()):
        if mgr.client:
            try:
                await mgr.client.disconnect()
            except Exception:
                pass

def create_app():
    app = web.Application(middlewares=[auth_middleware])
    app.on_startup.append(on_startup)
    app.on_cleanup.append(on_cleanup)

    # Main page & Static assets
    app.router.add_get("/", handle_index)
    app.router.add_static("/static/", path=STATIC_DIR, name="static")
    app.router.add_static("/uploads/", path=UPLOADS_DIR, name="uploads")

    # Auth routes
    app.router.add_post("/api/auth/register", handle_user_register)
    app.router.add_post("/api/auth/login", handle_user_login)
    app.router.add_post("/api/auth/google", handle_google_auth)
    app.router.add_post("/api/auth/logout", handle_user_logout)
    app.router.add_get("/api/auth/me", handle_user_me)
    app.router.add_get("/api/auth/firebase-config", handle_get_firebase_config)
    app.router.add_post("/api/auth/firebase-config", handle_save_firebase_config)

    # Telegram Account routes
    app.router.add_post("/api/auth/send_code", handle_auth_send_code)
    app.router.add_post("/api/auth/verify_code", handle_auth_verify_code)
    app.router.add_post("/api/auth/verify_password", handle_auth_verify_password)
    app.router.add_post("/api/auth/logout_telegram", handle_auth_logout_telegram)

    # User Settings & Status
    app.router.add_get("/api/status", handle_get_status)
    app.router.add_post("/api/config", handle_save_config)

    # File uploads & template
    app.router.add_post("/api/excel/upload", handle_upload_excel)
    app.router.add_post("/api/media/upload", handle_upload_media)
    app.router.add_get("/api/excel/template", handle_download_template)

    # Campaign routes
    app.router.add_post("/api/campaign/start", handle_campaign_start)
    app.router.add_post("/api/campaign/pause", handle_campaign_pause)
    app.router.add_post("/api/campaign/resume", handle_campaign_resume)
    app.router.add_post("/api/campaign/stop", handle_campaign_stop)
    app.router.add_get("/api/campaign/logs", handle_campaign_logs)
    app.router.add_get("/api/campaign/export", handle_export_excel)
    app.router.add_get("/api/campaign/history", handle_campaign_history)
    app.router.add_get("/api/campaign/history/{id}/download", handle_download_history_report)

    # WebSocket
    app.router.add_get("/ws", handle_ws)

    return app

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    host = os.environ.get("HOST", "0.0.0.0")
    print("=" * 60)
    print("[*] Telegram Excel Rassilka SaaS Platformasi ishga tushdi!")
    print(f"[*] Server manzili: http://{host}:{port}")
    print("=" * 60)

    if sys.platform == "win32" or os.environ.get("DISPLAY"):
        def open_browser():
            try:
                webbrowser.open(f"http://localhost:{port}")
            except Exception:
                pass
        threading.Timer(1.5, open_browser).start()

    app = create_app()
    web.run_app(app, host=host, port=port)
